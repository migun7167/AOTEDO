"""Report exports (FR-019): CSV, Excel, JSON and a raw-message package."""
from __future__ import annotations

import csv
import io
import json
import sqlite3
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AOT_NAVY = "0D2A5C"
AOT_GOLD = "C9A227"

SHEETS: dict[str, tuple[str, list[str]]] = {
    "Summary": ("""
        SELECT COALESCE(r.override_status, r.match_status) AS status,
               r.mawb_number, w.airline_prefix, w.origin, w.destination,
               w.flight_number, r.match_score, r.fhl_count,
               r.fwb_pieces, r.fhl_total_pieces, r.pieces_difference,
               r.fwb_weight, r.fhl_total_weight, r.weight_difference,
               r.reviewed, r.reviewed_by, r.review_note, r.last_matched_at
        FROM matching_results r
        LEFT JOIN fwb_master w ON w.id = r.fwb_id
        {where} ORDER BY r.mawb_number""", []),
    "FWB": ("""
        SELECT w.mawb_number, w.airline_prefix, w.origin, w.destination,
               w.pieces, w.gross_weight, w.weight_unit, w.chargeable_weight,
               w.flight_number, w.flight_date, w.shipper_name, w.consignee_name,
               w.agent_code, w.agent_name, w.currency, w.rate, w.freight_charge,
               w.total_charge, w.nature_of_goods, w.issue_date, w.issue_place,
               w.reference_number, w.created_at
        FROM fwb_master w
        JOIN matching_results r ON r.mawb_number = w.mawb_number
        {where} ORDER BY w.mawb_number""", []),
    "FHL": ("""
        SELECT h.mawb_number, h.hawb_number, h.origin, h.destination, h.pieces,
               h.gross_weight, h.weight_unit, h.commodity, h.hs_code,
               h.consignee_tax_id, h.shipper_name, h.consignee_name,
               h.consignee_country, h.consignee_phone, h.created_at
        FROM fhl_house h
        JOIN matching_results r ON r.mawb_number = h.mawb_number
        {where} ORDER BY h.mawb_number, h.hawb_number""", []),
    "Validation Results": ("""
        SELECT r.mawb_number, v.rule_code, v.severity, v.result,
               v.fwb_value, v.fhl_value, v.difference_value, v.created_at
        FROM validation_results v
        JOIN matching_results r ON r.id = v.matching_result_id
        {where} ORDER BY r.mawb_number, v.rule_code""", []),
}

ERRORS_SQL = """
    SELECT original_filename, message_type, message_version, parse_status,
           parse_error_code, parse_error_message, duplicate_type,
           imported_by, imported_at
    FROM cargo_messages
    WHERE parse_status IN ('PARSE_ERROR','INVALID_FORMAT','DUPLICATE')
       OR duplicate_type IS NOT NULL
    ORDER BY imported_at DESC"""

AUDIT_SQL = """
    SELECT event_type, user_id, entity_type, entity_id, reason, created_at
    FROM audit_logs ORDER BY created_at DESC LIMIT 5000"""


def build_match_where(status: str = "", date_from: str = "", date_to: str = "",
                      search: str = "") -> tuple[str, list]:
    """Filter clause shared by every export format (applies to matching_results)."""
    clauses, params = [], []
    if status:
        wanted = [s for s in status.split(",") if s]
        if wanted:
            clauses.append(
                "COALESCE(r.override_status, r.match_status) IN (%s)"
                % ",".join("?" * len(wanted)))
            params += wanted
    if date_from:
        clauses.append("substr(r.last_matched_at,1,10) >= ?")
        params.append(date_from)
    if date_to:
        clauses.append("substr(r.last_matched_at,1,10) <= ?")
        params.append(date_to)
    if search:
        clauses.append("r.mawb_number LIKE ?")
        params.append(f"%{search}%")
    return ("WHERE " + " AND ".join(clauses)) if clauses else "", params


def collect(conn: sqlite3.Connection, where: str,
            params: list) -> dict[str, list[dict]]:
    data: dict[str, list[dict]] = {}
    for name, (sql, _) in SHEETS.items():
        rows = conn.execute(sql.format(where=where), params).fetchall()
        data[name] = [dict(r) for r in rows]
    data["Errors"] = [dict(r) for r in conn.execute(ERRORS_SQL)]
    data["Audit Log"] = [dict(r) for r in conn.execute(AUDIT_SQL)]
    return data


def to_xlsx(data: dict[str, list[dict]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor=AOT_NAVY)
    header_font = Font(color="FFFFFF", bold=True, size=10)

    for sheet_name, rows in data.items():
        ws = wb.create_sheet(sheet_name[:31])
        if not rows:
            ws["A1"] = "ไม่มีข้อมูลตามเงื่อนไขที่เลือก"
            ws["A1"].font = Font(italic=True, color="888888")
            continue
        columns = list(rows[0].keys())
        ws.append(columns)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in rows:
            ws.append([row.get(c) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, column in enumerate(columns, start=1):
            widest = max([len(str(column))] +
                         [len(str(r.get(column) or "")) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(widest + 3, 46)

    meta = wb.create_sheet("Export Info")
    meta["A1"] = "Paperless AOT — FWB/FHL Matching Report"
    meta["A1"].font = Font(bold=True, size=14, color=AOT_NAVY)
    meta["A3"] = "Generated at"
    meta["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for i, (sheet_name, rows) in enumerate(data.items(), start=5):
        meta[f"A{i}"] = sheet_name
        meta[f"B{i}"] = len(rows)
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(data: dict[str, list[dict]]) -> bytes:
    """CSV carries the Summary sheet — the one row-per-MAWB view."""
    rows = data["Summary"]
    buf = io.StringIO()
    if rows:
        writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def to_json(data: dict[str, list[dict]]) -> bytes:
    payload = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sections": data,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def to_raw_package(conn: sqlite3.Connection, where: str, params: list) -> bytes:
    """Zip of the original message files, grouped by MAWB, plus a manifest."""
    mawbs = [r["mawb_number"] for r in conn.execute(
        f"SELECT r.mawb_number FROM matching_results r {where} ORDER BY 1", params)]
    buf = io.BytesIO()
    manifest = []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for mawb in mawbs:
            messages = conn.execute(
                """SELECT id, message_type, message_version, original_filename,
                          raw_message, imported_at
                   FROM cargo_messages WHERE id IN (
                     SELECT message_id FROM fwb_master WHERE mawb_number = ?
                     UNION SELECT message_id FROM fhl_house WHERE mawb_number = ?
                     UNION SELECT message_id FROM fsu_status WHERE mawb_number = ?
                     UNION SELECT message_id FROM ffm_flight WHERE mawb_number = ?)
                   ORDER BY imported_at""", (mawb,) * 4).fetchall()
            for i, m in enumerate(messages, start=1):
                name = m["original_filename"] or f"{m['message_type']}_{i}.txt"
                path = f"{mawb}/{i:02d}_{name}"
                zf.writestr(path, m["raw_message"])
                manifest.append({
                    "mawb": mawb, "file": path, "messageId": m["id"],
                    "type": m["message_type"], "version": m["message_version"],
                    "importedAt": m["imported_at"]})
        zf.writestr("manifest.json",
                    json.dumps(manifest, ensure_ascii=False, indent=2))
    return buf.getvalue()
